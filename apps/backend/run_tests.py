"""
Sahm Backend — Test Verification Runner
Runs validation, template engine, and batch tests using standard library unittest.
"""
import os
import sys
import unittest
import zipfile

# Add app directory to sys.path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass


from app.services.template_engine import TemplateFieldEngine, PRESET_TEMPLATES
from app.services.export_validator import ExportValidator
from app.services.renderers.pdf_renderer import PDFRenderer
from app.services.renderers.data_renderers import CSVRenderer, JSONRenderer, TXTRenderer
from app.services.renderers.batch_exporter import BatchExporter

try:
    import openpyxl
    from app.services.renderers.xlsx_renderer import XLSXRenderer
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class TestExportStudio(unittest.TestCase):

    def setUp(self):
        self.sample_records = [
            {
                "student_name": "أحمد عباس محمد إبراهيم",
                "university_id": "202201048",
                "status": "cert_ready",
                "specialization": "علوم الحاسوب",
                "phone": "+249912345678",
                "delivery_notes": "تسلم باليد بموجب التوكيل",
            },
            {
                "student_name": "محمد أحمد عثمان إدريس",
                "university_id": "002201052",
                "status": "approved",
                "specialization": "نظم المعلومات",
                "phone": "+249123456789",
                "delivery_notes": "",
            },
            {
                "student_name": "فاطمة الزهراء إدريس علي",
                "university_id": "202201079",
                "status": "cert_ready",
                "specialization": "تقنية المعلومات",
                "phone": "+249998877665",
                "delivery_notes": "مرتبة الشرف",
            },
        ]
        self.test_dir = os.path.join(os.path.dirname(__file__), "test_output")
        os.makedirs(self.test_dir, exist_ok=True)

    def test_01_template_field_engine(self):
        context = TemplateFieldEngine.build_context(
            university_name="جامعة إفريقيا العالمية",
            college_name="كلية دراسات الحاسوب",
            batch_year="2026",
            records_count=3,
        )
        template_str = "{{ university.name }} — {{ college.name }} ({{ batch.year }}) [{{ records.count }}]"
        resolved = TemplateFieldEngine.resolve_text(template_str, context)
        self.assertIn("جامعة إفريقيا العالمية", resolved)
        self.assertIn("كلية دراسات الحاسوب", resolved)
        self.assertIn("2026", resolved)
        self.assertIn("3", resolved)
        print("✓ Test 01 Passed: Template Field Engine token interpolation verified.")

    def test_02_export_validator_clean_data(self):
        report = ExportValidator.validate(
            records=self.sample_records,
            template_config=PRESET_TEMPLATES[0],
            is_public_publication=False,
        )
        self.assertTrue(report.is_valid)
        self.assertTrue(report.ready_to_export)
        self.assertEqual(len(report.errors), 0)
        self.assertIn("records_count_positive", report.passed_checks)
        print("✓ Test 02 Passed: Export Validator clean data verification passed.")

    def test_03_privacy_guard_blocks_phone(self):
        leaky_template = {
            "page_config": {"orientation": "portrait"},
            "table_config": {
                "columns": [
                    {"key": "student_name", "label": "الاسم", "visible": True},
                    {"key": "phone", "label": "الهاتف", "visible": True},
                ]
            },
        }
        report = ExportValidator.validate(
            records=self.sample_records,
            template_config=leaky_template,
            is_public_publication=True,
        )
        self.assertFalse(report.is_valid)
        self.assertFalse(report.ready_to_export)
        self.assertFalse(report.privacy_cleared)
        print("✓ Test 03 Passed: Privacy Guard successfully blocked restricted field 'phone' for public template.")

    def test_04_pdf_smart_pagination_renderer(self):
        renderer = PDFRenderer()
        out_file = os.path.join(self.test_dir, "test_doc.html")
        context = TemplateFieldEngine.build_context(
            records_count=3,
            additional_vars={"document.number": "CERT-2026-TEST", "verification.code": "VRF-UNITTEST-01"},
        )
        renderer.render(self.sample_records, PRESET_TEMPLATES[0], context, out_file)
        self.assertTrue(os.path.exists(out_file))

        with open(out_file, "r", encoding="utf-8") as f:
            html_text = f.read()

        self.assertIn('dir="rtl"', html_text)
        self.assertIn("IBM Plex Sans Arabic", html_text)
        self.assertIn("table-header-group", html_text)  # Repeated header rule
        self.assertIn("VRF-UNITTEST-01", html_text)      # QR Code badge
        print("✓ Test 04 Passed: PDF/HTML Smart Pagination & QR verification badge verified.")

    def test_05_csv_bom_encoding(self):
        renderer = CSVRenderer()
        out_csv = os.path.join(self.test_dir, "test.csv")
        renderer.render(self.sample_records, PRESET_TEMPLATES[0], {}, out_csv)
        self.assertTrue(os.path.exists(out_csv))

        # Check UTF-8 BOM
        with open(out_csv, "rb") as f:
            bom = f.read(3)
            self.assertEqual(bom, b"\xef\xbb\xbf")  # UTF-8 BOM for Excel Arabic
        print("✓ Test 05 Passed: CSV UTF-8 BOM encoding verified for seamless Excel Arabic opening.")

    def test_06_batch_zip_packaging(self):
        f1 = os.path.join(self.test_dir, "Report1.txt")
        with open(f1, "w", encoding="utf-8") as f:
            f.write("Batch Record 1")

        zip_out = os.path.join(self.test_dir, "Batch_Archive.zip")
        res = BatchExporter.create_batch_archive(
            generated_files=[{"status": "success", "entity_name": "كلية الحاسوب", "file_path": f1, "record_count": 50}],
            batch_name="Test_Batch_2026",
            output_zip_path=zip_out,
        )
        self.assertTrue(os.path.exists(zip_out))
        self.assertEqual(res["report"]["total_successful"], 1)

        with zipfile.ZipFile(zip_out, "r") as zf:
            self.assertIn("Report1.txt", zf.namelist())
            self.assertIn("BATCH_EXPORT_REPORT.json", zf.namelist())
        print("✓ Test 06 Passed: Batch packaging and manifest reporting verified.")

    def test_07_xlsx_renderer(self):
        if not OPENPYXL_AVAILABLE:
            print("⚠ Openpyxl not installed in current environment; skipping XLSX specific check.")
            return

        renderer = XLSXRenderer()
        out_xlsx = os.path.join(self.test_dir, "test.xlsx")
        context = TemplateFieldEngine.build_context(records_count=3)
        renderer.render(self.sample_records, PRESET_TEMPLATES[2], context, out_xlsx)
        self.assertTrue(os.path.exists(out_xlsx))

        wb = openpyxl.load_workbook(out_xlsx)
        self.assertIn("كشف الشهادات", wb.sheetnames)
        self.assertIn("الملخص الإحصائي", wb.sheetnames)
        ws = wb["كشف الشهادات"]
        self.assertTrue(ws.views.sheetView[0].rightToLeft)
        # Check string formatting for IDs
        id_cell = ws["C7"]
        self.assertEqual(id_cell.value, "002201052")
        print("✓ Test 07 Passed: Real XLSX with multiple sheets, RTL, and string IDs verified.")


if __name__ == "__main__":
    unittest.main()
